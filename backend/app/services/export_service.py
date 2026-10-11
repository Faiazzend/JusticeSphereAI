from datetime import datetime, timezone
from pathlib import Path
from typing import List, Dict, Any, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.config import settings


class ExportService:
    def __init__(self):
        pass

    def _add_track_change_deletion(
        self, paragraph, text: str, change_id: int, utc_iso: str
    ):
        del_elem = OxmlElement("w:del")
        del_elem.set(qn("w:id"), str(change_id))
        del_elem.set(qn("w:author"), "ContractMind Legal AI")
        del_elem.set(qn("w:date"), utc_iso)

        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")

        strike = OxmlElement("w:strike")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "DC2626")

        rPr.append(strike)
        rPr.append(color)
        r.append(rPr)

        del_text = OxmlElement("w:delText")
        del_text.text = text
        del_text.set(qn("xml:space"), "preserve")
        r.append(del_text)

        del_elem.append(r)
        paragraph._p.append(del_elem)

    def _add_track_change_insertion(
        self, paragraph, text: str, change_id: int, utc_iso: str
    ):
        ins_elem = OxmlElement("w:ins")
        ins_elem.set(qn("w:id"), str(change_id))
        ins_elem.set(qn("w:author"), "ContractMind Legal AI")
        ins_elem.set(qn("w:date"), utc_iso)

        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")

        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "16A34A")

        rPr.append(u)
        rPr.append(color)
        r.append(rPr)

        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)

        ins_elem.append(r)
        paragraph._p.append(ins_elem)

    def generate_docx_redline(
        self,
        contract_title: str,
        contract_id: str,
        overall_risk_score: float,
        risk_summary: Dict[str, int],
        redlines: List[Dict[str, Any]],
        obligations: Optional[List[Dict[str, Any]]] = None,
    ) -> str:
        """
        Creates a native Word document containing authentic OpenXML Track Changes
        insertions and deletions with decision branching (Accepted, Rejected, Pending Review),
        an executive summary table, and legal analysis.
        """
        doc = Document()
        utc_iso = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        change_counter = 1

        # Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("ContractMind v1.5 — Legal Audit & Redline Synthesis")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)

        subtitle_p = doc.add_paragraph()
        sub_run = subtitle_p.add_run(
            f"Document: {contract_title} | ID: {contract_id} | Generated: {utc_iso}"
        )
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph()

        # Section 1: Executive Risk Summary
        h1 = doc.add_heading("1. Executive Risk & Scorecard Summary", level=1)
        h1.paragraph_format.space_before = Pt(12)

        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        metrics = [
            ("Composite Risk Score", f"{overall_risk_score} / 100"),
            ("Critical Risk Count", str(risk_summary.get("critical", 0))),
            ("High Risk Count", str(risk_summary.get("high", 0))),
            ("Medium Risk Count", str(risk_summary.get("medium", 0))),
            ("Low Risk Count", str(risk_summary.get("low", 0))),
        ]

        for i, (label, val) in enumerate(metrics):
            row = table.rows[i]
            row.cells[0].paragraphs[0].add_run(label).bold = True
            row.cells[1].paragraphs[0].add_run(val)

        doc.add_paragraph()

        # Section 2: Redline Decisions Overview Table
        h2 = doc.add_heading("2. Redline Decisions Overview", level=1)
        h2.paragraph_format.space_before = Pt(14)

        if redlines:
            ov_table = doc.add_table(rows=1 + len(redlines), cols=4)
            ov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ov_table.style = "Table Grid"

            headers = ["Section", "Severity", "Redline Type", "Decision Status"]
            for col_idx, header in enumerate(headers):
                cell = ov_table.rows[0].cells[col_idx]
                r = cell.paragraphs[0].add_run(header)
                r.bold = True

            for row_idx, item in enumerate(redlines):
                sec_num = item.get("section_number") or item.get("clause_id") or f"Section {row_idx+1}"
                sev = item.get("risk_severity") or "HIGH"
                rtype = item.get("diff_type") or item.get("category") or "Modification"
                accepted = item.get("accepted")

                if accepted is True:
                    status_text = "Accepted"
                elif accepted is False:
                    status_text = "Rejected"
                else:
                    status_text = "Pending Review"

                row = ov_table.rows[row_idx + 1]
                row.cells[0].paragraphs[0].add_run(str(sec_num))
                row.cells[1].paragraphs[0].add_run(str(sev))
                row.cells[2].paragraphs[0].add_run(str(rtype))
                r_status = row.cells[3].paragraphs[0].add_run(status_text)
                if status_text == "Accepted":
                    r_status.font.color.rgb = RGBColor(22, 163, 74)
                elif status_text == "Rejected":
                    r_status.font.color.rgb = RGBColor(220, 38, 38)
                else:
                    r_status.font.color.rgb = RGBColor(79, 70, 229)

        doc.add_paragraph()

        # Section 3: Detailed Clauses with Track Changes and Decision Branching
        h3 = doc.add_heading("3. Detailed Clause Redlines (OpenXML Track Changes)", level=1)
        h3.paragraph_format.space_before = Pt(14)

        if not redlines:
            p = doc.add_paragraph("No contractual redlines flagged for this document.")
            p.runs[0].font.italic = True
        else:
            for idx, item in enumerate(redlines):
                sec_num = item.get("section_number") or item.get("clause_id") or f"Section {idx+1}"
                sev = item.get("risk_severity") or "HIGH"
                cat = item.get("diff_type") or item.get("category") or "Uncategorized"
                accepted = item.get("accepted")
                diff_tokens = item.get("diff_tokens", [])
                orig_text = item.get("original_text", "")
                rationale = item.get("legal_rationale", "Standard consensus compromise.")
                custom_prompt = item.get("custom_prompt")

                if accepted is True:
                    status_text = "ACCEPTED"
                elif accepted is False:
                    status_text = "REJECTED"
                else:
                    status_text = "PENDING REVIEW"

                # Heading incorporates decision status
                clause_h = doc.add_heading(
                    f"3.{idx+1} {sec_num} [Severity: {sev} | Decision: {status_text}]",
                    level=2,
                )
                clause_h.paragraph_format.space_before = Pt(10)

                # Decision Branching Output:
                if accepted is False:
                    # REJECTED: Original clause retained, no Track Changes insertions
                    dec_p = doc.add_paragraph()
                    r_dec = dec_p.add_run("Decision: REJECTED (Original Clause Retained): ")
                    r_dec.bold = True
                    r_dec.font.color.rgb = RGBColor(220, 38, 38)

                    orig_p = doc.add_paragraph()
                    orig_p.paragraph_format.left_indent = Inches(0.2)
                    orig_p.add_run(orig_text)

                else:
                    # ACCEPTED or PENDING: Inject genuine Track Changes OpenXML
                    dec_p = doc.add_paragraph()
                    r_dec = dec_p.add_run(f"Decision: {status_text} (Track Changes XML Enabled): ")
                    r_dec.bold = True
                    if accepted is True:
                        r_dec.font.color.rgb = RGBColor(22, 163, 74)
                    else:
                        r_dec.font.color.rgb = RGBColor(79, 70, 229)

                    tc_para = doc.add_paragraph()
                    tc_para.paragraph_format.left_indent = Inches(0.2)

                    for token in diff_tokens:
                        token_type = token.get("type", "equal")
                        token_text = token.get("text", "")
                        if not token_text:
                            continue

                        if token_type == "equal":
                            tc_para.add_run(token_text)
                        elif token_type == "delete":
                            self._add_track_change_deletion(
                                tc_para, token_text, change_counter, utc_iso
                            )
                            change_counter += 1
                        elif token_type == "insert":
                            self._add_track_change_insertion(
                                tc_para, token_text, change_counter, utc_iso
                            )
                            change_counter += 1

                # Custom Prompt callout if present
                if custom_prompt:
                    cp_para = doc.add_paragraph()
                    cp_run = cp_para.add_run("Lawyer Custom Prompt: ")
                    cp_run.bold = True
                    cp_para.add_run(custom_prompt)
                    cp_para.paragraph_format.left_indent = Inches(0.2)

                # Legal Rationale Box
                rat_para = doc.add_paragraph()
                rat_run = rat_para.add_run("Legal Rationale & Context: ")
                rat_run.bold = True
                rat_para.add_run(rationale)
                rat_para.paragraph_format.space_after = Pt(8)

        # Section 4: Obligations Matrix
        if obligations:
            h4 = doc.add_heading("4. Structured Obligations & Compliance Calendar", level=1)
            h4.paragraph_format.space_before = Pt(14)

            ob_table = doc.add_table(rows=1 + len(obligations), cols=5)
            ob_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ob_table.style = "Table Grid"

            headers = ["Type", "Obligor", "Recipient", "Due Date / Trigger", "Penalty for Breach"]
            for col_idx, header in enumerate(headers):
                cell = ob_table.rows[0].cells[col_idx]
                r = cell.paragraphs[0].add_run(header)
                r.bold = True

            for row_idx, ob in enumerate(obligations):
                row = ob_table.rows[row_idx + 1]
                row.cells[0].paragraphs[0].add_run(str(ob.get("obligation_type", "")))
                row.cells[1].paragraphs[0].add_run(str(ob.get("obligor", "")))
                row.cells[2].paragraphs[0].add_run(str(ob.get("recipient", "")))
                row.cells[3].paragraphs[0].add_run(str(ob.get("due_date_or_trigger", "")))
                row.cells[4].paragraphs[0].add_run(str(ob.get("penalty_for_breach", "")))

        # Save document
        filename = f"ContractMind_Redline_{contract_id[:8]}.docx"
        out_path = settings.EXPORT_DIR / filename
        doc.save(str(out_path))
        return str(out_path)


export_service = ExportService()
