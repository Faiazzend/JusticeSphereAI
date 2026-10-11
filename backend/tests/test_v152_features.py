import pytest
import pytest_asyncio
import httpx
from docx import Document
from pathlib import Path

from app.main import app
from app.database import init_db
from app.services.export_service import export_service


@pytest.mark.asyncio
async def test_v152_interactive_redlines_and_filters():
    await init_db()

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Ingest benchmark contract
        res = await client.post("/api/v1/contracts/demo/ma")
        assert res.status_code == 200
        contract = res.json()
        cid = contract["id"]
        redlines = contract["redlines"]
        assert len(redlines) > 0
        rl_id = redlines[0]["id"]

        # 2. Test GET /analysis/{cid}/redlines without filter
        res = await client.get(f"/api/v1/analysis/{cid}/redlines")
        assert res.status_code == 200
        assert len(res.json()) == len(redlines)

        # 3. Test GET /analysis/{cid}/redlines with risk_severity filter
        res = await client.get(f"/api/v1/analysis/{cid}/redlines?risk_severity=CRITICAL")
        assert res.status_code == 200
        critical_rls = res.json()
        assert len(critical_rls) > 0
        for rl in critical_rls:
            assert rl["risk_severity"] == "CRITICAL"

        # 4. Test GET /analysis/{cid}/redlines with decision filter
        res = await client.get(f"/api/v1/analysis/{cid}/redlines?decision=pending")
        assert res.status_code == 200
        assert len(res.json()) == len(redlines)

        # 5. Test PATCH /analysis/{cid}/redlines/{rl_id} - ACCEPT
        res = await client.patch(
            f"/api/v1/analysis/{cid}/redlines/{rl_id}",
            json={"accepted": True},
        )
        assert res.status_code == 200
        updated_rl = res.json()
        assert updated_rl["accepted"] is True

        # Verify filter by accepted decision
        res = await client.get(f"/api/v1/analysis/{cid}/redlines?decision=accepted")
        assert res.status_code == 200
        assert len(res.json()) == 1
        assert res.json()[0]["id"] == rl_id

        # 6. Test PATCH - REJECT
        res = await client.patch(
            f"/api/v1/analysis/{cid}/redlines/{rl_id}",
            json={"accepted": False},
        )
        assert res.status_code == 200
        assert res.json()["accepted"] is False

        # Verify filter by rejected decision
        res = await client.get(f"/api/v1/analysis/{cid}/redlines?decision=rejected")
        assert res.status_code == 200
        assert len(res.json()) == 1

        # 7. Test POST /analysis/{cid}/redlines/{rl_id}/regenerate
        custom_prompt = "Make indemnification strictly mutual and cap at $500,000."
        res = await client.post(
            f"/api/v1/analysis/{cid}/redlines/{rl_id}/regenerate",
            json={"custom_prompt": custom_prompt},
        )
        assert res.status_code == 200
        regen_rl = res.json()
        assert regen_rl["custom_prompt"] == custom_prompt
        assert regen_rl["accepted"] is None  # Reset to pending review
        assert len(regen_rl["diff_tokens"]) > 0

        # Verify contract transcript updated
        res = await client.get(f"/api/v1/contracts/{cid}")
        contract_data = res.json()
        transcript = contract_data["debate_transcript"]
        assert any(t["agent_name"] == "Client Prompt & Arbiter" for t in transcript)


def test_docx_decision_branching():
    redlines = [
        {
            "id": "rl_1",
            "section_number": "Section 2.1",
            "risk_severity": "CRITICAL",
            "diff_type": "Indemnification",
            "original_text": "Seller provides uncapped first-dollar tipping indemnity.",
            "suggested_text": "Seller provides deductible basket capped at 1.0x fees.",
            "diff_tokens": [
                {"type": "equal", "text": "Seller provides "},
                {"type": "delete", "text": "uncapped first-dollar tipping indemnity."},
                {"type": "insert", "text": "deductible basket capped at 1.0x fees."},
            ],
            "accepted": True,
            "legal_rationale": "Accepted revision.",
        },
        {
            "id": "rl_2",
            "section_number": "Section 3.1",
            "risk_severity": "HIGH",
            "diff_type": "Warranties",
            "original_text": "Representations made without inquiry.",
            "suggested_text": "Representations made after reasonable inquiry.",
            "diff_tokens": [
                {"type": "delete", "text": "without inquiry."},
                {"type": "insert", "text": "after reasonable inquiry."},
            ],
            "accepted": False,  # REJECTED
            "legal_rationale": "Rejected revision.",
        },
    ]

    out_file = export_service.generate_docx_redline(
        contract_title="Decision Branching Contract",
        contract_id="test-branch-uuid",
        overall_risk_score=65.0,
        risk_summary={"critical": 1, "high": 1, "medium": 0, "low": 0},
        redlines=redlines,
    )

    assert Path(out_file).exists()
    doc = Document(out_file)
    xml_content = doc._element.xml

    # Check Overview Table headers
    assert any("Redline Decisions Overview" in p.text for p in doc.paragraphs)
    # Check decision text
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Decision: ACCEPTED (Track Changes XML Enabled):" in full_text
    assert "Decision: REJECTED (Original Clause Retained):" in full_text
    assert "Representations made without inquiry." in full_text
    assert "w:ins" in xml_content
    assert "w:del" in xml_content
